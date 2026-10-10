from pathlib import Path

import pytest

import config
from core.document_loader import Block, extract_text, load_document


@pytest.fixture(autouse=True)
def _ocr_off_by_default(monkeypatch):
    """P2.2 的按页 Block 用例默认关闭 OCR，避免扫描件判定介入改变 Block 文本；
    需要验证 OCR 行为的用例（P2.5）在自身内重新开启 ``config.OCR_ENABLED``。"""
    monkeypatch.setattr(config, "OCR_ENABLED", False)


def _build_pdf(page_texts):
    """手工构造一个最小多页 PDF（仅用内置 Helvetica，不引入第三方写库）。

    ``page_texts`` 为每页文本（空串表示空白页）。返回 PDF 字节串。
    """
    header = b"%PDF-1.4\n"
    objects = {}
    n_pages = len(page_texts)
    first_page_obj = 4  # 1=Catalog, 2=Pages, 3=Font

    page_obj_nums = [first_page_obj + 2 * i for i in range(n_pages)]
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = b" ".join(b"%d 0 R" % n for n in page_obj_nums)
    objects[2] = b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % n_pages
    objects[3] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    for i, text in enumerate(page_texts):
        page_num = page_obj_nums[i]
        content_num = page_num + 1
        if text:
            stream = ("BT /F1 12 Tf 72 720 Td (%s) Tj ET" % text).encode("latin-1")
        else:
            stream = b""
        objects[page_num] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % content_num
        )
        objects[content_num] = (
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        )

    out = bytearray(header)
    offsets = {}
    for num in sorted(objects):
        offsets[num] = len(out)
        out += b"%d 0 obj\n" % num + objects[num] + b"\nendobj\n"

    xref_offset = len(out)
    max_num = max(objects) + 1
    out += b"xref\n0 %d\n" % max_num
    out += b"0000000000 65535 f \n"
    for num in range(1, max_num):
        out += b"%010d 00000 n \n" % offsets[num]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        max_num,
        xref_offset,
    )
    return bytes(out)


def test_extract_utf8_text_and_markdown(tmp_path: Path):
    text_file = tmp_path / "sample.txt"
    text_file.write_text("RAG combines retrieval and generation.\n中文内容。", encoding="utf-8")

    markdown_file = tmp_path / "sample.md"
    markdown_file.write_text("# Heading\n\nHybrid retrieval", encoding="utf-8")

    assert "中文内容" in extract_text(str(text_file))
    assert "Hybrid retrieval" in extract_text(str(markdown_file))


def test_unsupported_extension_returns_empty_string(tmp_path: Path):
    unsupported = tmp_path / "sample.bin"
    unsupported.write_bytes(b"not a supported document")

    assert extract_text(str(unsupported)) == ""


def test_extract_docx_paragraphs(tmp_path: Path):
    from docx import Document

    docx_file = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("RAG docx paragraph")
    document.save(str(docx_file))

    assert "RAG docx paragraph" in extract_text(str(docx_file))


def test_extract_xlsx_sheets(tmp_path: Path):
    import pandas as pd

    xlsx_file = tmp_path / "sample.xlsx"
    pd.DataFrame({"col": ["value1", "value2"]}).to_excel(xlsx_file, index=False, sheet_name="Sheet1")

    text = extract_text(str(xlsx_file))
    assert "Sheet1" in text
    assert "value1" in text


def test_extract_pptx_text(tmp_path: Path):
    from pptx import Presentation

    pptx_file = tmp_path / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "RAG pptx title"
    presentation.save(str(pptx_file))

    assert "RAG pptx title" in extract_text(str(pptx_file))


def test_load_document_pdf_pages_start_at_one(tmp_path: Path):
    pdf_file = tmp_path / "multi.pdf"
    pdf_file.write_bytes(_build_pdf(["ALPHA_PAGE_ONE", "BRAVO_PAGE_TWO", "CHARLIE_PAGE_THREE"]))

    blocks = load_document(str(pdf_file))

    assert all(isinstance(block, Block) for block in blocks)
    assert [block.page for block in blocks] == [1, 2, 3]
    assert all(block.source == "multi.pdf" for block in blocks)
    assert all(block.ocr_used is False for block in blocks)


def test_load_document_pdf_pages_are_independent(tmp_path: Path):
    pdf_file = tmp_path / "multi.pdf"
    pdf_file.write_bytes(_build_pdf(["ALPHA_PAGE_ONE", "BRAVO_PAGE_TWO", "CHARLIE_PAGE_THREE"]))

    blocks = load_document(str(pdf_file))

    assert "ALPHA_PAGE_ONE" in blocks[0].text
    assert "BRAVO_PAGE_TWO" not in blocks[0].text
    assert "BRAVO_PAGE_TWO" in blocks[1].text
    assert "ALPHA_PAGE_ONE" not in blocks[1].text
    assert "CHARLIE_PAGE_THREE" in blocks[2].text


def test_load_document_pdf_blank_page_keeps_real_page_count(tmp_path: Path):
    pdf_file = tmp_path / "with_blank.pdf"
    pdf_file.write_bytes(_build_pdf(["ONLY_ONE", "", "THREE_HERE"]))

    blocks = load_document(str(pdf_file))

    assert [block.page for block in blocks] == [1, 2, 3]
    assert blocks[1].text.strip() == ""
    assert "ONLY_ONE" in blocks[0].text
    assert "THREE_HERE" in blocks[2].text


def test_load_document_non_pdf_returns_single_page_zero_block(tmp_path: Path):
    text_file = tmp_path / "sample.txt"
    text_file.write_text("RAG 单页文本", encoding="utf-8")

    blocks = load_document(str(text_file))

    assert len(blocks) == 1
    assert blocks[0].page == 0
    assert blocks[0].source == "sample.txt"
    assert blocks[0].ocr_used is False
    assert "RAG 单页文本" in blocks[0].text


def test_load_document_unsupported_extension_returns_empty_block(tmp_path: Path):
    unsupported = tmp_path / "sample.bin"
    unsupported.write_bytes(b"not a supported document")

    blocks = load_document(str(unsupported))

    assert len(blocks) == 1
    assert blocks[0].text == ""
    assert blocks[0].page == 0
    assert blocks[0].source == "sample.bin"


def test_extract_text_still_wraps_pdf_pages(tmp_path: Path):
    pdf_file = tmp_path / "multi.pdf"
    pdf_file.write_bytes(_build_pdf(["ALPHA_PAGE_ONE", "BRAVO_PAGE_TWO"]))

    text = extract_text(str(pdf_file))

    assert "ALPHA_PAGE_ONE" in text
    assert "BRAVO_PAGE_TWO" in text


def test_load_document_text_pdf_does_not_trigger_ocr(tmp_path: Path, monkeypatch):
    import core.ocr as ocr

    monkeypatch.setattr(config, "OCR_ENABLED", True)
    calls = {"render": 0}
    monkeypatch.setattr(
        ocr, "render_pdf_pages", lambda path: calls.__setitem__("render", calls["render"] + 1) or []
    )

    pdf_file = tmp_path / "text.pdf"
    pdf_file.write_bytes(_build_pdf(["A" * 400]))

    blocks = load_document(str(pdf_file))

    assert calls["render"] == 0
    assert all(block.ocr_used is False for block in blocks)


def test_load_document_scanned_pdf_runs_ocr(tmp_path: Path, monkeypatch):
    import numpy as np
    import core.ocr as ocr

    monkeypatch.setattr(config, "OCR_ENABLED", True)
    monkeypatch.setattr(
        ocr, "render_pdf_pages",
        lambda path: [ocr.RenderedPage(page=1, status="ok", image=np.zeros((4, 4, 3), dtype="uint8"))],
    )
    monkeypatch.setattr(ocr, "ocr_image", lambda image: ocr.OcrResult(text="OCR 文本", lines=[]))

    pdf_file = tmp_path / "scan.pdf"
    pdf_file.write_bytes(_build_pdf([""]))  # 无文本层 → 判定为扫描件

    blocks = load_document(str(pdf_file))

    assert blocks[0].ocr_used is True
    assert blocks[0].text == "OCR 文本"


def test_load_document_low_text_page_triggers_page_level_ocr(tmp_path: Path, monkeypatch):
    import numpy as np
    import core.ocr as ocr

    monkeypatch.setattr(config, "OCR_ENABLED", True)
    rendered = [
        ocr.RenderedPage(page=i, status="ok", image=np.zeros((4, 4, 3), dtype="uint8"))
        for i in (1, 2, 3)
    ]
    monkeypatch.setattr(ocr, "render_pdf_pages", lambda path: rendered)
    ocr_calls = []
    monkeypatch.setattr(
        ocr, "ocr_image", lambda image: ocr_calls.append(True) or ocr.OcrResult(text="补足", lines=[])
    )

    pdf_file = tmp_path / "mixed.pdf"
    pdf_file.write_bytes(_build_pdf(["P" * 400, "", "Q" * 400]))  # 第 2 页为低文本页

    blocks = load_document(str(pdf_file))

    assert [block.ocr_used for block in blocks] == [False, True, False]
    assert blocks[1].text == "补足"
    assert len(ocr_calls) == 1


def test_load_document_force_ocr_overrides_detection(tmp_path: Path, monkeypatch):
    import numpy as np
    import core.ocr as ocr

    monkeypatch.setattr(config, "OCR_ENABLED", True)
    monkeypatch.setattr(
        ocr, "render_pdf_pages",
        lambda path: [ocr.RenderedPage(page=1, status="ok", image=np.zeros((4, 4, 3), dtype="uint8"))],
    )
    monkeypatch.setattr(ocr, "ocr_image", lambda image: ocr.OcrResult(text="强制", lines=[]))

    pdf_file = tmp_path / "text.pdf"
    pdf_file.write_bytes(_build_pdf(["A" * 400]))

    blocks = load_document(str(pdf_file), force_ocr=True)

    assert blocks[0].ocr_used is True
    assert blocks[0].text == "强制"
